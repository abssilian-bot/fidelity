# -*- coding: utf-8 -*-
"""Modèle financier 12 mois — 3 scénarios (pessimiste / réaliste / ambitieux).
On mange quoi / Fidelity. Tout est piloté par les hypothèses (cellules bleues).
Le script écrit les formules ET injecte les valeurs calculées en cache."""
import zipfile, shutil, re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.utils import get_column_letter

OUT = r"C:\Users\jebza\Documents\Kimi\Workspaces\OnMangeQuoi\modele-financier-12mois.xlsx"

BLUE, GREEN, BLACK, GRAY, LIGHT = "0066CC", "1E7B34", "000000", "666666", "F5F5F5"
MONEY, MONEY0, PCT = '#,##0.00" €"', '#,##0" €"', '0%'

caches = {}  # sheet_name -> {cellref: value}

# ------------------------------------------------------------------ moteur de calcul
def compute(p, restos, conv, wl):
    """p = dict d'hypothèses. Retourne {row: [12 valeurs]} + totaux."""
    rows = {r: [] for r in (22, 23, 25, 26, 28, 30, 31, 32, 33, 34, 36, 37)}
    cash = 0.0
    for m in range(12):
        premium = restos[m] * conv[m]
        rec = restos[m] * p["rec_resto"]
        ca_abo = premium * p["prix"]
        ca_rec = rec * p["frais"]
        ca = ca_abo + ca_rec + wl[m]
        c_infra = p["infra"] if (m + 1) >= p["infra_mois"] else 0
        c_apple = p["apple"] if (m + 1) >= p["apple_mois"] else 0
        c_stripe = ca_abo * p["stripe_pct"] + premium * p["stripe_fixe"]
        couts = c_infra + c_apple + p["domaine"] + c_stripe
        net = ca - couts
        cash += net
        for r, v in [(22, premium), (23, rec), (25, ca_abo), (26, ca_rec), (28, ca),
                     (30, c_infra), (31, c_apple), (32, p["domaine"]), (33, c_stripe),
                     (34, couts), (36, net), (37, cash)]:
            rows[r].append(v)
    return rows

# ------------------------------------------------------------------ construction d'une feuille scénario
def build_scenario(wb, name, p, restos, conv, wl, notes):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 50
    for j in range(5, 18):
        ws.column_dimensions[get_column_letter(j)].width = 11

    ws.row_dimensions[2].height = 30
    ws.merge_cells("B2:D2")
    ws["B2"] = f"Scénario {name.lower()} — hypothèses et projection 12 mois"
    ws["B2"].font = Font(size=16, bold=True)

    # ---- bloc hypothèses (lignes 4-14), références en $C$5..$C$14
    for j, h in enumerate(["Paramètre", "Valeur", "Note"]):
        c = ws.cell(row=4, column=2 + j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="333333")
    params = [
        ("Prix abonnement premium (€/mois)", p["prix"], MONEY0, notes["prix"]),
        ("Frais par récompense débloquée (€)", p["frais"], MONEY, "Prélevé uniquement quand un client utilise sa récompense"),
        ("Récompenses débloquées / restaurant actif / mois", p["rec_resto"], "0", notes["rec"]),
        ("Commission Stripe variable (%)", p["stripe_pct"], '0.0%', "Sur les abonnements encaissés"),
        ("Commission Stripe fixe (€/abonné/mois)", p["stripe_fixe"], MONEY, "Par restaurant premium facturé"),
        ("Nom de domaine (€/mois)", p["domaine"], MONEY, "~10 €/an lissés"),
        ("Apple Developer (€/mois)", p["apple"], MONEY, "99 $/an ≈ 93 €/an, lissés"),
        ("Mois d'activation Apple Developer", p["apple_mois"], "0", notes["apple"]),
        ("Infra au-delà des free tiers (€/mois)", p["infra"], MONEY0, notes["infra"]),
        ("Mois de bascule infra payante", p["infra_mois"], "0", notes["infra_mois"]),
    ]
    for i, (label, val, fmt, note) in enumerate(params):
        r = 5 + i
        ws.cell(row=r, column=2, value=label).font = Font(size=10)
        c = ws.cell(row=r, column=3, value=val)
        c.font = Font(size=10, bold=True, color=BLUE)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4, value=note).font = Font(size=9, color=GRAY)

    H = {k: f"$C${5 + i}" for i, k in enumerate(
        ["prix", "frais", "rec_resto", "stripe_pct", "stripe_fixe",
         "domaine", "apple", "apple_mois", "infra", "infra_mois"])}

    # ---- modèle : mois en colonnes E..P (décalé car le bloc hypo occupe B-D)
    base_col = 5  # E
    total_col = 17  # Q
    ws.cell(row=16, column=2, value="Projection mensuelle").font = Font(size=13, bold=True)

    ws.cell(row=17, column=4, value="N° du mois").font = Font(size=8, color=GRAY)
    ws.cell(row=18, column=4, value="Ligne").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=18, column=4).fill = PatternFill("solid", start_color="333333")
    ws.cell(row=18, column=total_col, value="Total").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=18, column=total_col).fill = PatternFill("solid", start_color="333333")
    ws.cell(row=18, column=total_col).alignment = Alignment(horizontal="center")
    for m in range(1, 13):
        col = base_col + m - 1
        ci = ws.cell(row=17, column=col, value=m)
        ci.font = Font(size=8, color=GRAY)
        ci.alignment = Alignment(horizontal="center")
        ch = ws.cell(row=18, column=col, value=f"M{m}")
        ch.font = Font(bold=True, color="FFFFFF")
        ch.fill = PatternFill("solid", start_color="333333")
        ch.alignment = Alignment(horizontal="center")

    cache = {}

    def section(row, label):
        c = ws.cell(row=row, column=4, value=label)
        c.font = Font(bold=True, size=11)
        for j in range(4, total_col + 1):
            ws.cell(row=row, column=j).fill = PatternFill("solid", start_color=LIGHT)

    def row_label(row, label, bold=False):
        ws.cell(row=row, column=4, value=label).font = Font(size=11, bold=bold)

    def fill_row(row, values=None, formula=None, fmt=None, blue=False, bold=False,
                 total=None, computed=None):
        for m in range(1, 13):
            col = base_col + m - 1
            L = get_column_letter(col)
            if values is not None:
                c = ws.cell(row=row, column=col, value=values[m - 1])
                c.font = Font(size=11, color=BLUE if blue else BLACK, bold=bold)
            else:
                c = ws.cell(row=row, column=col, value=formula(L))
                c.font = Font(size=11, color=BLACK, bold=bold)
                cache[f"{L}{row}"] = computed[m - 1]
            if fmt:
                c.number_format = fmt
            c.alignment = Alignment(horizontal="center")
        if total is not None:
            c = ws.cell(row=row, column=total_col, value=total)
            c.font = Font(size=11, bold=True, color=BLACK)
            c.number_format = fmt or "General"
            c.alignment = Alignment(horizontal="center")

    rows = compute(p, restos, conv, wl)
    TL = get_column_letter(total_col)  # Q

    section(19, "VOLUME")
    row_label(20, "Restaurants actifs (saisie)")
    fill_row(20, values=restos, fmt="0", blue=True, total="=P20")
    cache[f"{TL}20"] = restos[-1]
    row_label(21, "Taux de conversion premium (saisie)")
    fill_row(21, values=conv, fmt=PCT, blue=True)
    row_label(22, "Restaurants premium")
    fill_row(22, formula=lambda L: f"={L}20*{L}21", fmt="0.0", computed=rows[22])
    row_label(23, "Récompenses débloquées dans le mois")
    fill_row(23, formula=lambda L: f"={L}20*{H['rec_resto']}", fmt="0",
             total="=SUM(E23:P23)", computed=rows[23])
    cache[f"{TL}23"] = sum(rows[23])

    section(24, "REVENUS (€)")
    row_label(25, "Abonnements premium")
    fill_row(25, formula=lambda L: f"={L}22*{H['prix']}", fmt=MONEY, computed=rows[25],
             total="=SUM(E25:P25)")
    cache[f"{TL}25"] = sum(rows[25])
    row_label(26, "Frais sur récompenses débloquées")
    fill_row(26, formula=lambda L: f"={L}23*{H['frais']}", fmt=MONEY, computed=rows[26],
             total="=SUM(E26:P26)")
    cache[f"{TL}26"] = sum(rows[26])
    row_label(27, "White-label / pub (saisie)")
    fill_row(27, values=wl, fmt=MONEY0, blue=True, total="=SUM(E27:P27)")
    cache[f"{TL}27"] = sum(wl)
    row_label(28, "CA total", bold=True)
    fill_row(28, formula=lambda L: f"=SUM({L}25:{L}27)", fmt=MONEY, bold=True,
             computed=rows[28], total="=SUM(E28:P28)")
    cache[f"{TL}28"] = sum(rows[28])
    for j in range(4, total_col + 1):
        ws.cell(row=28, column=j).border = Border(top=Side(style="thin", color=BLACK),
                                                  bottom=Side(style="double", color=BLACK))

    section(29, "COÛTS (€)")
    row_label(30, "Infra (hébergement, BDD, stockage)")
    fill_row(30, formula=lambda L: f"=IF({L}$17>={H['infra_mois']},{H['infra']},0)",
             fmt=MONEY, computed=rows[30], total="=SUM(E30:P30)")
    cache[f"{TL}30"] = sum(rows[30])
    row_label(31, "Apple Developer")
    fill_row(31, formula=lambda L: f"=IF({L}$17>={H['apple_mois']},{H['apple']},0)",
             fmt=MONEY, computed=rows[31], total="=SUM(E31:P31)")
    cache[f"{TL}31"] = sum(rows[31])
    row_label(32, "Nom de domaine")
    fill_row(32, formula=lambda L: f"={H['domaine']}", fmt=MONEY, computed=rows[32],
             total="=SUM(E32:P32)")
    cache[f"{TL}32"] = sum(rows[32])
    row_label(33, "Commissions Stripe")
    fill_row(33, formula=lambda L: f"={L}25*{H['stripe_pct']}+{L}22*{H['stripe_fixe']}",
             fmt=MONEY, computed=rows[33], total="=SUM(E33:P33)")
    cache[f"{TL}33"] = sum(rows[33])
    row_label(34, "Coûts totaux", bold=True)
    fill_row(34, formula=lambda L: f"=SUM({L}30:{L}33)", fmt=MONEY, bold=True,
             computed=rows[34], total="=SUM(E34:P34)")
    cache[f"{TL}34"] = sum(rows[34])

    section(35, "RÉSULTAT")
    row_label(36, "Résultat net du mois", bold=True)
    fill_row(36, formula=lambda L: f"={L}28-{L}34", fmt=MONEY, bold=True,
             computed=rows[36], total="=SUM(E36:P36)")
    cache[f"{TL}36"] = sum(rows[36])
    row_label(37, "Trésorerie cumulée")
    first_col = get_column_letter(base_col)
    c = ws.cell(row=37, column=base_col, value=f"={first_col}36")
    c.font = Font(size=11); c.number_format = MONEY; c.alignment = Alignment(horizontal="center")
    cache[f"{first_col}37"] = rows[37][0]
    for m in range(2, 13):
        col = base_col + m - 1
        L, P = get_column_letter(col), get_column_letter(col - 1)
        c = ws.cell(row=37, column=col, value=f"={P}37+{L}36")
        c.font = Font(size=11); c.number_format = MONEY; c.alignment = Alignment(horizontal="center")
        cache[f"{L}37"] = rows[37][m - 1]

    # ---- graphique
    chart = LineChart()
    chart.title = "CA total, coûts et trésorerie cumulée (€)"
    chart.style = 2
    chart.height = 9
    chart.width = 26
    for rr in (28, 34, 37):
        ref = Reference(ws, min_col=4, min_row=rr, max_col=base_col + 11, max_row=rr)
        chart.add_data(ref, titles_from_data=True, from_rows=True)
    cats = Reference(ws, min_col=base_col, min_row=18, max_col=base_col + 11, max_row=18)
    chart.set_categories(cats)
    chart.y_axis.title = "€"
    chart.x_axis.title = "Mois"
    ws.add_chart(chart, "B40")

    caches[name] = cache
    return ws, rows

# ------------------------------------------------------------------ paramètres des 3 scénarios
P_PESS = dict(prix=19, frais=0.30, rec_resto=6, stripe_pct=0.015, stripe_fixe=0.25,
              domaine=0.83, apple=7.75, apple_mois=4, infra=22, infra_mois=9)
R_PESS = [1, 1, 2, 2, 3, 4, 4, 5, 6, 7, 8, 9]
C_PESS = [0, 0, 0, .05, .05, .08, .08, .10, .10, .12, .12, .15]
W_PESS = [0] * 12

P_REAL = dict(prix=24, frais=0.30, rec_resto=12, stripe_pct=0.015, stripe_fixe=0.25,
              domaine=0.83, apple=7.75, apple_mois=3, infra=30, infra_mois=6)
R_REAL = [1, 2, 3, 5, 7, 10, 13, 16, 20, 25, 30, 38]
C_REAL = [0, 0, .05, .08, .10, .12, .15, .18, .20, .22, .25, .28]
W_REAL = [0, 0, 0, 0, 0, 0, 0, 0, 0, 500, 500, 500]

P_AMB = dict(prix=29, frais=0.30, rec_resto=20, stripe_pct=0.015, stripe_fixe=0.25,
             domaine=0.83, apple=7.75, apple_mois=2, infra=45, infra_mois=4)
R_AMB = [2, 4, 7, 12, 18, 26, 36, 48, 62, 80, 100, 125]
C_AMB = [0, .05, .08, .12, .15, .18, .22, .25, .28, .32, .35, .38]
W_AMB = [0, 0, 0, 0, 0, 500, 1000, 1000, 1500, 1500, 2000, 2000]

wb = Workbook()

# ------------------------------------------------------------------ Couverture
cov = wb.active
cov.title = "Couverture"
cov.sheet_view.showGridLines = False
cov.column_dimensions["A"].width = 2
cov.column_dimensions["B"].width = 30
for col, w in (("C", 16), ("D", 16), ("E", 18), ("F", 55)):
    cov.column_dimensions[col].width = w

cov.row_dimensions[2].height = 34
cov.merge_cells("B2:F2")
cov["B2"] = "On mange quoi / Fidelity — Modèle financier 12 mois, 3 scénarios"
cov["B2"].font = Font(size=18, bold=True)
cov.merge_cells("B3:F3")
cov["B3"] = "Pessimiste · Réaliste · Ambitieux — 31 août 2026"
cov["B3"].font = Font(size=11, color=GRAY)

cov["B5"] = "Vue d'ensemble — année 1"
cov["B5"].font = Font(size=13, bold=True)
for j, h in enumerate(["Scénario", "CA année", "Résultat net", "Trésorerie fin M12"]):
    c = cov.cell(row=6, column=2 + j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="333333")
    c.alignment = Alignment(horizontal="center" if j else "left")

# references filled after scenario sheets exist (cell map known upfront)
scen_refs = [("Pessimiste (gratter le fond)", "Modèle", 28, 36, 37),
             ("Réaliste", "Réaliste", 28, 36, 37),
             ("Ambitieux", "Ambitieux", 28, 36, 37)]
# NB : pour 'Modèle' (pessimiste), les lignes seront celles de la feuille reconstruite ci-dessous
cover_cache = {}
for i, (label, sheet, r_ca, r_net, r_cash) in enumerate(scen_refs):
    r = 7 + i
    cov.cell(row=r, column=2, value=label).font = Font(size=11)
    for j, (col, rr) in enumerate((("Q", r_ca), ("Q", r_net), ("P", r_cash))):
        c = cov.cell(row=r, column=3 + j, value=f"='{sheet}'!{col}{rr}")
        c.font = Font(size=11, bold=True, color=GREEN)
        c.number_format = MONEY0
        c.alignment = Alignment(horizontal="center")

cov["B11"] = "Détails fin M12"
cov["B11"].font = Font(size=13, bold=True)
for j, h in enumerate(["Scénario", "Restos actifs", "Restos premium", "Mois bénéficiaires"]):
    c = cov.cell(row=12, column=2 + j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="333333")
for i, (label, sheet, *_rest) in enumerate(scen_refs):
    r = 13 + i
    cov.cell(row=r, column=2, value=label).font = Font(size=11)
    c = cov.cell(row=r, column=3, value=f"='{sheet}'!P20")
    c.font = Font(size=11, color=GREEN); c.number_format = "0"; c.alignment = Alignment(horizontal="center")
    c = cov.cell(row=r, column=4, value=f"='{sheet}'!P22")
    c.font = Font(size=11, color=GREEN); c.number_format = "0.0"; c.alignment = Alignment(horizontal="center")
    c = cov.cell(row=r, column=5, value=f"=COUNTIF('{sheet}'!E36:P36,\">0\")&\" / 12\"")
    c.font = Font(size=11, color=GREEN); c.alignment = Alignment(horizontal="center")

cov["B17"] = "Contenu du classeur"
cov["B17"].font = Font(size=13, bold=True)
index = [
    ("Modèle", "Scénario pessimiste : croissance minimale, 9 restos fin M12, aucun white-label"),
    ("Réaliste", "Bouche-à-oreille correct dans une ville : 38 restos fin M12, 1 contrat white-label"),
    ("Ambitieux", "Ça prend : 125 restos fin M12, white-label + pub dès le 2e semestre"),
]
r = 18
for name, desc in index:
    cov.cell(row=r, column=2, value=name).font = Font(size=11, bold=True)
    cov.cell(row=r, column=3, value=desc).font = Font(size=10, color=GRAY)
    cov.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    r += 1

cov["B22"] = "Mode d'emploi"
cov["B22"].font = Font(size=13, bold=True)
notes_cov = [
    "Chaque scénario est autonome : ses hypothèses (cellules bleues) sont en haut de sa feuille.",
    "Modifiez une cellule bleue, toute la feuille se recalcule.",
    "Convention couleurs : bleu = saisie, noir = calcul, vert = référence à une autre feuille.",
]
r = 23
for n in notes_cov:
    cov.cell(row=r, column=2, value=n).font = Font(size=10, color=GRAY)
    cov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    r += 1

# ------------------------------------------------------------------ feuilles scénarios
notes_pess = {"prix": "Fourchette basse", "rec": "Base clients très petite",
              "apple": "Phase 2 Wallet", "infra": "Free tiers longtemps suffisants",
              "infra_mois": "Croissance lente = bascule tardive"}
notes_real = {"prix": "Milieu de fourchette 19-29 €", "rec": "Base clients en croissance",
              "apple": "Wallet activé tôt", "infra": "Dépassement free tiers vers ~500 utilisateurs actifs",
              "infra_mois": "La courbe d'usage le justifie"}
notes_amb = {"prix": "Haut de fourchette, services inclus (photoshoot…)",
             "rec": "Usage intensif, FOODSHARE actif", "apple": "Wallet dès le début",
             "infra": "Pic d'usage + stockage images", "infra_mois": "Trafic précoce"}

mod_ws, rows_pess = build_scenario(wb, "Modèle", P_PESS, R_PESS, C_PESS, W_PESS, notes_pess)
real_ws, rows_real = build_scenario(wb, "Réaliste", P_REAL, R_REAL, C_REAL, W_REAL, notes_real)
amb_ws, rows_amb = build_scenario(wb, "Ambitieux", P_AMB, R_AMB, C_AMB, W_AMB, notes_amb)

# title cell of pessimiste sheet mentions its name properly
mod_ws["B2"] = "Scénario pessimiste — hypothèses et projection 12 mois"

# ------------------------------------------------------------------ caches couverture
scen_rows = {"Modèle": rows_pess, "Réaliste": rows_real, "Ambitieux": rows_amb}
scen_restos = {"Modèle": R_PESS, "Réaliste": R_REAL, "Ambitieux": R_AMB}
for i, (label, sheet, *_r) in enumerate(scen_refs):
    rows = scen_rows[sheet]
    cover_cache[f"C{7 + i}"] = sum(rows[28])
    cover_cache[f"D{7 + i}"] = sum(rows[36])
    cover_cache[f"E{7 + i}"] = rows[37][-1]
    cover_cache[f"C{13 + i}"] = scen_restos[sheet][-1]
    cover_cache[f"D{13 + i}"] = rows[22][-1]
    cover_cache[f"E{13 + i}"] = f"{sum(1 for v in rows[36] if v > 0)} / 12"
caches["Couverture"] = cover_cache

# ------------------------------------------------------------------ Comparaison
cmp_ws = wb.create_sheet("Comparaison")
cmp_ws.sheet_view.showGridLines = False
cmp_ws.column_dimensions["A"].width = 2
cmp_ws.column_dimensions["B"].width = 26
for col in "CDEFG":
    cmp_ws.column_dimensions[col].width = 17

cmp_ws.row_dimensions[2].height = 30
cmp_ws.merge_cells("B2:G2")
cmp_ws["B2"] = "Comparaison des 3 scénarios — année 1"
cmp_ws["B2"].font = Font(size=16, bold=True)

headers = ["Scénario", "CA année", "Coûts année", "Résultat net", "CA mensuel M12", "Restos premium M12"]
for j, h in enumerate(headers):
    c = cmp_ws.cell(row=4, column=2 + j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="333333")
    c.alignment = Alignment(horizontal="center" if j else "left", wrap_text=True)
cmp_ws.row_dimensions[4].height = 28

cmp_cache = {}
for i, (label, sheet, *_r) in enumerate(scen_refs):
    r = 5 + i
    rows = scen_rows[sheet]
    cmp_ws.cell(row=r, column=2, value=label).font = Font(size=11)
    specs = [("Q", 28, MONEY0), ("Q", 34, MONEY0), ("Q", 36, MONEY0), ("P", 28, MONEY0), ("P", 22, "0.0")]
    vals = [sum(rows[28]), sum(rows[34]), sum(rows[36]), rows[28][-1], rows[22][-1]]
    for j, ((cl, rr, fmt), v) in enumerate(zip(specs, vals)):
        c = cmp_ws.cell(row=r, column=3 + j, value=f"='{sheet}'!{cl}{rr}")
        c.font = Font(size=11, color=GREEN)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="center")
        cmp_cache[f"{get_column_letter(3 + j)}{r}"] = v
caches["Comparaison"] = cmp_cache

bchart = BarChart()
bchart.type = "col"
bchart.title = "CA année vs résultat net par scénario (€)"
bchart.style = 2
bchart.height = 9
bchart.width = 22
data = Reference(cmp_ws, min_col=3, min_row=4, max_col=3, max_row=7)
bchart.add_data(data, titles_from_data=True)
data2 = Reference(cmp_ws, min_col=5, min_row=4, max_col=5, max_row=7)
bchart.add_data(data2, titles_from_data=True)
cats = Reference(cmp_ws, min_col=2, min_row=5, max_row=7)
bchart.set_categories(cats)
bchart.y_axis.title = "€"
cmp_ws.add_chart(bchart, "B10")

# ------------------------------------------------------------------ sauvegarde + injection des caches
wb.save(OUT)

zin = zipfile.ZipFile(OUT)
wbxml = zin.read("xl/workbook.xml").decode("utf-8")
rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
sheet_ids = re.findall(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="(rId\d+)"', wbxml)
rel_map = {}
for m in re.finditer(r'<Relationship [^>]*?Target="/?xl/(worksheets/[^"]+)"[^>]*?Id="(rId\d+)"', rels):
    rel_map[m.group(2)] = m.group(1)
for m in re.finditer(r'<Relationship [^>]*?Id="(rId\d+)"[^>]*?Target="/?xl/(worksheets/[^"]+)"', rels):
    rel_map.setdefault(m.group(1), m.group(2))
name_to_file = {name: "xl/" + rel_map[rid] for name, rid in sheet_ids if rid in rel_map}
print("mapping:", name_to_file)

targets = {name_to_file[name]: cellvals for name, cellvals in caches.items() if name in name_to_file}

def inject(xml, cellvals):
    for ref, v in cellvals.items():
        pat = re.compile(r'<c r="%s"[^>]*>(<f>.*?</f>)' % ref)
        m = pat.search(xml)
        if not m:
            print("!! cellule non trouvée:", ref)
            continue
        full = m.group(0)
        if isinstance(v, str):
            src = full if ' t=' in full[:full.index('>')] else full.replace('<c r=', '<c t="str" r=', 1)
            new = re.sub(r'(<f>.*?</f>)', lambda mm: mm.group(1) + f'<v>{v}</v>', src)
        else:
            new = re.sub(r'(<f>.*?</f>)', lambda mm: mm.group(1) + f'<v>{round(float(v), 10)}</v>', full)
        xml = xml[:m.start()] + new + xml[m.end():]
    return xml

TMP = OUT + ".tmp"
with zipfile.ZipFile(TMP, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in targets:
            data = inject(data.decode("utf-8"), targets[item.filename]).encode("utf-8")
        zout.writestr(item, data)
zin.close()
shutil.move(TMP, OUT)
print("OK:", OUT)

# ------------------------------------------------------------------ résumé console
for label, sheet, *_r in scen_refs:
    rows = scen_rows[sheet]
    print(f"\n=== {sheet} ===")
    print(f"CA année        : {sum(rows[28]):>10,.2f} €")
    print(f"Coûts année     : {sum(rows[34]):>10,.2f} €")
    print(f"Résultat net    : {sum(rows[36]):>10,.2f} €")
    print(f"CA mensuel M12  : {rows[28][-1]:>10,.2f} €")
    print(f"Restos premium M12 : {rows[22][-1]:.1f}")
