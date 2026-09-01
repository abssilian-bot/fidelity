# -*- coding: utf-8 -*-
"""Modèle financier 12 mois — scénario volontairement pessimiste.
On mange quoi / Fidelity. Tout est piloté par les hypothèses (cellules bleues)."""
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

OUT = r"C:\Users\jebza\Documents\Kimi\Workspaces\OnMangeQuoi\modele-financier-12mois.xlsx"

BLUE = "0066CC"      # inputs
GREEN = "1E7B34"     # cross-sheet refs
BLACK = "000000"
GRAY = "666666"
LIGHT = "F5F5F5"
MONEY = '#,##0.00" €"'
MONEY0 = '#,##0" €"'
PCT = '0%'

wb = Workbook()

# ---------------------------------------------------------------- Couverture
cov = wb.active
cov.title = "Couverture"
cov.sheet_view.showGridLines = False
cov.column_dimensions["A"].width = 2
cov.column_dimensions["B"].width = 46
cov.column_dimensions["C"].width = 18
cov.column_dimensions["D"].width = 60

cov.row_dimensions[2].height = 34
cov.merge_cells("B2:D2")
cov["B2"] = "On mange quoi / Fidelity — Modèle financier 12 mois"
cov["B2"].font = Font(size=18, bold=True, color=BLACK)

cov.merge_cells("B3:D3")
cov["B3"] = "Scénario volontairement pessimiste · croissance très lente · 31 août 2026"
cov["B3"].font = Font(size=11, color=GRAY)

cov["B5"] = "Indicateurs clés (année 1)"
cov["B5"].font = Font(size=13, bold=True)
metrics = [
    ("CA total année", "='Modèle'!O14", MONEY),
    ("Coûts totaux année", "='Modèle'!O20", MONEY),
    ("Résultat net année", "='Modèle'!O22", MONEY),
    ("Trésorerie cumulée fin M12", "='Modèle'!N23", MONEY),
    ("Restaurants actifs fin M12", "='Modèle'!N6", "0"),
    ("Mois bénéficiaires (résultat > 0)", '=COUNTIF(\'Modèle\'!C22:N22,">0")&" / 12"', "@"),
]
r = 6
for label, formula, fmt in metrics:
    cov.cell(row=r, column=2, value=label).font = Font(size=11)
    c = cov.cell(row=r, column=3, value=formula)
    c.font = Font(size=11, bold=True, color=GREEN)
    if fmt != "@":
        c.number_format = fmt
    r += 1

cov["B13"] = "Contenu du classeur"
cov["B13"].font = Font(size=13, bold=True)
index = [
    ("Hypothèses", "Tous les paramètres modifiables (cellules bleues) : prix, frais, coûts, calendrier"),
    ("Modèle", "Projection mensuelle M1 → M12 : volumes, revenus, coûts, résultat, trésorerie + graphique"),
]
r = 14
for name, desc in index:
    cov.cell(row=r, column=2, value=name).font = Font(size=11, bold=True)
    cov.cell(row=r, column=3, value=desc).font = Font(size=10, color=GRAY)
    cov.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    r += 1

cov["B17"] = "Mode d'emploi"
cov["B17"].font = Font(size=13, bold=True)
notes = [
    "Les cellules bleues sont des hypothèses : modifiez-les, tout le modèle se recalcule.",
    "Hypothèses volontairement défavorables : démarrage à 1 restaurant, 9 restaurants fin M12,",
    "conversion premium plafonnée à 15 %, aucun contrat white-label sur l'année.",
    "Convention couleurs : bleu = saisie, noir = calcul, vert = référence à une autre feuille.",
]
r = 18
for n in notes:
    cov.cell(row=r, column=2, value=n).font = Font(size=10, color=GRAY)
    cov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1

# ---------------------------------------------------------------- Hypothèses
hyp = wb.create_sheet("Hypothèses")
hyp.sheet_view.showGridLines = False
hyp.column_dimensions["A"].width = 2
hyp.column_dimensions["B"].width = 52
hyp.column_dimensions["C"].width = 14
hyp.column_dimensions["D"].width = 58

hyp.row_dimensions[2].height = 30
hyp.merge_cells("B2:D2")
hyp["B2"] = "Hypothèses — scénario pessimiste"
hyp["B2"].font = Font(size=16, bold=True)

hdr = ["Paramètre", "Valeur", "Note"]
for j, h in enumerate(hdr):
    c = hyp.cell(row=4, column=2 + j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", start_color="333333")

params = [
    # label, value, fmt, note
    ("Prix abonnement premium (€/mois)", 19, MONEY0, "Fourchette basse 19-29 € retenue au plus bas"),
    ("Frais par récompense débloquée (€)", 0.30, MONEY, "Prélevé uniquement quand un client utilise sa récompense"),
    ("Récompenses débloquées / restaurant actif / mois", 6, "0", "Très faible : base clients petite la 1re année"),
    ("Commission Stripe variable (%)", 0.015, '0.0%', "Sur les abonnements encaissés"),
    ("Commission Stripe fixe (€/abonné/mois)", 0.25, MONEY, "Par restaurant premium facturé"),
    ("Nom de domaine (€/mois)", 0.83, MONEY, "~10 €/an lissés"),
    ("Apple Developer (€/mois)", 7.75, MONEY, "99 $/an ≈ 93 €/an, lissés"),
    ("Mois d'activation Apple Developer", 4, "0", "Phase 2 Wallet (feuille de route)"),
    ("Infra au-delà des free tiers (€/mois)", 22, MONEY0, "Supabase/Render payants (~500+ utilisateurs actifs)"),
    ("Mois de bascule infra payante", 9, "0", "Croissance lente = bascule tardive"),
]
r = 5
for label, val, fmt, note in params:
    hyp.cell(row=r, column=2, value=label).font = Font(size=11)
    c = hyp.cell(row=r, column=3, value=val)
    c.font = Font(size=11, bold=True, color=BLUE)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="center")
    hyp.cell(row=r, column=4, value=note).font = Font(size=10, color=GRAY)
    r += 1

# cell refs used by the model
H = {"prix": "Hypothèses!$C$5", "frais": "Hypothèses!$C$6",
     "rec_par_resto": "Hypothèses!$C$7", "stripe_pct": "Hypothèses!$C$8",
     "stripe_fixe": "Hypothèses!$C$9", "domaine": "Hypothèses!$C$10",
     "apple": "Hypothèses!$C$11", "apple_mois": "Hypothèses!$C$12",
     "infra": "Hypothèses!$C$13", "infra_mois": "Hypothèses!$C$14"}

# ---------------------------------------------------------------- Modèle
mod = wb.create_sheet("Modèle")
mod.sheet_view.showGridLines = False
mod.column_dimensions["A"].width = 2
mod.column_dimensions["B"].width = 40
for j in range(3, 16):  # C..O
    mod.column_dimensions[get_column_letter(j)].width = 11

mod.row_dimensions[2].height = 30
mod.merge_cells("B2:O2")
mod["B2"] = "Projection mensuelle — 12 mois"
mod["B2"].font = Font(size=16, bold=True)

# month index row 3 (helper) + header row 4
mod.cell(row=3, column=2, value="N° du mois").font = Font(size=8, color=GRAY)
mod.cell(row=4, column=2, value="Ligne").font = Font(bold=True, color="FFFFFF")
mod.cell(row=4, column=2).fill = PatternFill("solid", start_color="333333")
mod.cell(row=4, column=15, value="Total").font = Font(bold=True, color="FFFFFF")
mod.cell(row=4, column=15).fill = PatternFill("solid", start_color="333333")
mod.cell(row=4, column=15).alignment = Alignment(horizontal="center")
for m in range(1, 13):
    col = 2 + m
    ci = mod.cell(row=3, column=col, value=m)
    ci.font = Font(size=8, color=GRAY)
    ci.alignment = Alignment(horizontal="center")
    ch = mod.cell(row=4, column=col, value=f"M{m}")
    ch.font = Font(bold=True, color="FFFFFF")
    ch.fill = PatternFill("solid", start_color="333333")
    ch.alignment = Alignment(horizontal="center")

thin = Border(bottom=Side(style="thin", color="D0D0D0"))

def section(row, label):
    c = mod.cell(row=row, column=2, value=label)
    c.font = Font(bold=True, size=11)
    for j in range(2, 16):
        mod.cell(row=row, column=j).fill = PatternFill("solid", start_color=LIGHT)

def row_label(row, label, bold=False):
    mod.cell(row=row, column=2, value=label).font = Font(size=11, bold=bold)

def fill_row(row, values=None, formula=None, fmt=None, blue=False, bold=False,
             total=None, total_fmt=None):
    """values: list of 12 static inputs; formula: callable(col_letter)->str."""
    for m in range(1, 13):
        col = 2 + m
        letter = get_column_letter(col)
        if values is not None:
            c = mod.cell(row=row, column=col, value=values[m - 1])
            c.font = Font(size=11, color=BLUE if blue else BLACK, bold=bold)
        else:
            c = mod.cell(row=row, column=col, value=formula(letter))
            c.font = Font(size=11, color=BLACK, bold=bold)
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal="center")
    if total is not None:
        c = mod.cell(row=row, column=15, value=total)
        c.font = Font(size=11, bold=True, color=BLACK)
        c.number_format = total_fmt or fmt or "General"
        c.alignment = Alignment(horizontal="center")

# ---- VOLUME
section(5, "VOLUME")
row_label(6, "Restaurants actifs (saisie)")
fill_row(6, values=[1, 1, 2, 2, 3, 4, 4, 5, 6, 7, 8, 9], fmt="0", blue=True,
         total="=N6", total_fmt="0")
row_label(7, "Taux de conversion premium (saisie)")
fill_row(7, values=[0, 0, 0, 0.05, 0.05, 0.08, 0.08, 0.10, 0.10, 0.12, 0.12, 0.15],
         fmt=PCT, blue=True)
row_label(8, "Restaurants premium")
fill_row(8, formula=lambda L: f"={L}6*{L}7", fmt="0.0")
row_label(9, "Récompenses débloquées dans le mois")
fill_row(9, formula=lambda L: f"={L}6*{H['rec_par_resto']}", fmt="0",
         total="=SUM(C9:N9)", total_fmt="0")

# ---- REVENUS
section(10, "REVENUS (€)")
row_label(11, "Abonnements premium")
fill_row(11, formula=lambda L: f"={L}8*{H['prix']}", fmt=MONEY, total="=SUM(C11:N11)")
row_label(12, "Frais sur récompenses débloquées")
fill_row(12, formula=lambda L: f"={L}9*{H['frais']}", fmt=MONEY, total="=SUM(C12:N12)")
row_label(13, "White-label (saisie)")
fill_row(13, values=[0] * 12, fmt=MONEY, blue=True, total="=SUM(C13:N13)")
row_label(14, "CA total", bold=True)
fill_row(14, formula=lambda L: f"=SUM({L}11:{L}13)", fmt=MONEY, bold=True,
         total="=SUM(C14:N14)")
for j in range(2, 16):
    mod.cell(row=14, column=j).border = Border(top=Side(style="thin", color=BLACK),
                                               bottom=Side(style="double", color=BLACK))

# ---- COÛTS
section(15, "COÛTS (€)")
row_label(16, "Infra (hébergement, BDD, stockage)")
fill_row(16, formula=lambda L: f"=IF({L}$3>={H['infra_mois']},{H['infra']},0)",
         fmt=MONEY, total="=SUM(C16:N16)")
row_label(17, "Apple Developer")
fill_row(17, formula=lambda L: f"=IF({L}$3>={H['apple_mois']},{H['apple']},0)",
         fmt=MONEY, total="=SUM(C17:N17)")
row_label(18, "Nom de domaine")
fill_row(18, formula=lambda L: f"={H['domaine']}", fmt=MONEY, total="=SUM(C18:N18)")
row_label(19, "Commissions Stripe")
fill_row(19, formula=lambda L: f"={L}11*{H['stripe_pct']}+{L}8*{H['stripe_fixe']}",
         fmt=MONEY, total="=SUM(C19:N19)")
row_label(20, "Coûts totaux", bold=True)
fill_row(20, formula=lambda L: f"=SUM({L}16:{L}19)", fmt=MONEY, bold=True,
         total="=SUM(C20:N20)")

# ---- RÉSULTAT
section(21, "RÉSULTAT")
row_label(22, "Résultat net du mois", bold=True)
fill_row(22, formula=lambda L: f"={L}14-{L}20", fmt=MONEY, bold=True,
         total="=SUM(C22:N22)")
row_label(23, "Trésorerie cumulée")
mod.cell(row=23, column=3, value="=C22").font = Font(size=11)
mod.cell(row=23, column=3).number_format = MONEY
mod.cell(row=23, column=3).alignment = Alignment(horizontal="center")
for m in range(2, 13):
    col = 2 + m
    L = get_column_letter(col)
    P = get_column_letter(col - 1)
    c = mod.cell(row=23, column=col, value=f"={P}23+{L}22")
    c.font = Font(size=11)
    c.number_format = MONEY
    c.alignment = Alignment(horizontal="center")

# ---- Chart
chart = LineChart()
chart.title = "CA total, coûts et trésorerie cumulée (€)"
chart.style = 2
chart.height = 9
chart.width = 26
data = Reference(mod, min_col=2, min_row=14, max_col=14, max_row=14)  # CA
chart.add_data(data, titles_from_data=True, from_rows=True)
data2 = Reference(mod, min_col=2, min_row=20, max_col=14, max_row=20)
chart.add_data(data2, titles_from_data=True, from_rows=True)
data3 = Reference(mod, min_col=2, min_row=23, max_col=14, max_row=23)
chart.add_data(data3, titles_from_data=True, from_rows=True)
cats = Reference(mod, min_col=3, min_row=4, max_col=14, max_row=4)
chart.set_categories(cats)
chart.y_axis.title = "€"
chart.x_axis.title = "Mois"
mod.add_chart(chart, "B26")

wb.save(OUT)
print("OK:", OUT)
